"""Track B: ingestion of platform CSV/Excel exports (spec candidate-discovery).

Humans export from LinkedIn Recruiter / Naukri ResDEX UIs and drop files into
exports-inbox/. This module maps known column layouts into the normalized
candidate schema - it NEVER talks to those platforms itself.

Format registry is data-driven: each entry matches on a required-column
signature. New export formats are added by appending entries (task 6.1 grows
this table as real stakeholder exports are observed).
"""
from __future__ import annotations

import csv
from pathlib import Path

from candidate_schema import make_record


class UnknownFormatError(ValueError):
    pass


# --------------------------------------------------------------- normalization

def _norm(s: str) -> str:
    return (s or "").strip().lower().replace(" ", "_").replace(".", "")


def _to_years(value: str):
    if not value:
        return None
    digits = "".join(ch for ch in value if ch.isdigit())
    return int(digits) if digits else None


# ------------------------------------------------------------------- formats
# PROVISIONAL signatures - validated against real exports in task 6.1.

FORMATS = [
    {
        "name": "linkedin-recruiter",
        "required": {"first_name", "current_title", "profile_url"},
        "map": lambda row: make_record(
            name=" ".join(x for x in (row.get("first_name"),
                                      row.get("last_name")) if x).strip() or None,
            title=row.get("current_title"),
            company=row.get("current_company"),
            location=row.get("location"),
            experience_years=None,
            track="platform-export",
            source_urls=[row["profile_url"]] if row.get("profile_url") else [],
        ),
    },
    {
        "name": "naukri-resdex",
        "required": {"name", "resume_title", "resume_id"},
        "map": lambda row: make_record(
            name=row.get("name"),
            title=row.get("resume_title"),
            company=row.get("current_employer"),
            location=row.get("location"),
            experience_years=_to_years(row.get("total_experience")),
            track="platform-export",
            source_urls=[f"naukri:{row['resume_id']}"],
        ),
    },
]


def detect_format(header: list[str]) -> dict | None:
    cols = {_norm(h) for h in header}
    for fmt in FORMATS:
        if fmt["required"] <= cols:
            return fmt
    return None


def ingest_bytes(name: str, content: bytes) -> tuple[list[dict], dict | None]:
    """Parse one exported file -> (records, format_info|None).

    Raises UnknownFormatError for unrecognized layouts (caller quarantines).
    """
    suffix = Path(name).suffix.lower()
    if suffix == ".csv":
        rows = _read_csv(content)
    elif suffix in (".xlsx", ".xls"):
        rows = _read_excel(content)
    else:
        raise UnknownFormatError(f"unsupported file type: {name}")

    if not rows:
        raise UnknownFormatError(f"no rows found in {name}")
    header = list(rows[0].keys())
    fmt = detect_format(header)
    if fmt is None:
        raise UnknownFormatError(
            f"{name}: no format signature matches columns {sorted(header)}"
        )
    return [fmt["map"](row) for row in rows], {"format": fmt["name"], "rows": len(rows)}


def _read_csv(content: bytes) -> list[dict]:
    import io
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{_norm(k): v for k, v in row.items()} for row in reader]


def _read_excel(content: bytes) -> list[dict]:
    import io

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    values = ws.iter_rows(values_only=True)
    try:
        header = [str(c) if c is not None else "" for c in next(values)]
    except StopIteration:
        return []
    rows = []
    for raw in values:
        row = {_norm(h): v for h, v in zip(header, raw) if h}
        if any(v is not None for v in row.values()):
            rows.append(row)
    return rows


def ingest_inbox(list_names_fn, read_fn, mark_processed_fn, quarantine_fn) -> dict:
    """Process every file in exports-inbox (move-after-success semantics).

    Spec scenarios implemented here:
      - recognized exports -> records tagged platform-export, file archived
      - unrecognized/corrupt file -> quarantined to jds-failed, named in summary,
        other inbox files still processed
    Returns a run-summary dict: {'ingested': [...], 'quarantined': [...]}.
    """
    summary = {"ingested": [], "quarantined": []}
    for name in sorted(list_names_fn()):
        try:
            content = read_fn(name)
            records, info = ingest_bytes(name, content)
        except Exception as exc:  # corrupt or unrecognized -> quarantine, continue
            summary["quarantined"].append({"file": name, "reason": str(exc)})
            quarantine_fn(name)
            continue
        summary["ingested"].append({"file": name, "records": records, **info})
        mark_processed_fn(name)
    return summary
