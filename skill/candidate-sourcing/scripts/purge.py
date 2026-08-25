"""Retention purge helper (task 10.3, config: retention.purge_rejected_after_months).

Removes REJECTED ledger rows older than the configured window.
Dry-run is the default; --apply performs the deletion. Other statuses
(contacted/interviewing/hired) are never purged by this tool.
"""
from __future__ import annotations

import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from config_loader import load_config


def _parse_iso(value) -> date | None:
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None


def plan_purge(path: Path, cfg: dict | None = None,
               today: date | None = None) -> list[int]:
    """Row indexes (1-based worksheet rows) eligible for purge."""
    cfg = cfg or load_config()
    months = int(cfg.get("retention", {}).get("purge_rejected_after_months", 6))
    cutoff = (today or date.today()) - timedelta(days=30 * months)

    ws = load_workbook(path).active
    header = [str(c or "").lower() for c in next(ws.iter_rows(values_only=True))]
    try:
        col_date = header.index("date_added")
        col_status = header.index("status")
    except ValueError as exc:
        raise SystemExit(f"ledger missing required columns: {exc}")

    doomed = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if idx == 1:
            continue
        status = str(row[col_status] or "").strip().lower()
        added = _parse_iso(row[col_date])
        if status == "rejected" and added and added < cutoff:
            doomed.append(idx)
    return doomed


def apply_purge(path: Path, row_indexes: list[int]) -> int:
    ws = load_workbook(path).active
    for idx in sorted(row_indexes, reverse=True):   # bottom-up keeps indexes valid
        ws.delete_rows(idx)
    ws.parent.save(path)
    return len(row_indexes)


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 1
    path = Path(argv[1])
    doomed = plan_purge(path)
    if not doomed:
        print("nothing to purge")
        return 0
    print(f"{len(doomed)} rejected row(s) older than retention window:")
    for idx in doomed:
        print(f"  worksheet row {idx}")
    if "--apply" in argv:
        apply_purge(path, doomed)
        print("deleted.")
    else:
        print("dry-run only; re-run with --apply to delete")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
