"""Excel shortlist ledger (spec: shortlist-ledger).

The workbook is BOTH the human deliverable and the workflow's dedup memory:
append-only candidate rows, reviewer-controlled `status` column driving
exclusions. Guardrails per spec: backup before every write, refuse-and-defer
when the file looks locked (Excel lock file / permission error).
"""
from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

LEDGER_COLUMNS = [
    "date_added", "jd_title", "source_jd", "candidate_name", "title",
    "current_company", "location", "track", "profile_url", "source_urls",
    "score", "justification", "status", "possible_duplicate_flag", "notes",
]

TERMINAL_STATUSES = {"rejected", "contacted", "interviewing", "hired"}


class LedgerLockedError(RuntimeError):
    """Raised when the workbook appears open elsewhere -> defer to next run."""


def _lock_files(path: Path) -> list[Path]:
    return list(path.parent.glob(f"~${path.name}"))


def ensure_workbook(path: Path) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "shortlist"
    ws.append(LEDGER_COLUMNS)
    from openpyxl.styles import Font
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(path)


def _record_to_row(record: dict, jd_title: str,
                   source_jd: str | None = None) -> list:
    return [
        date.today().isoformat(),
        jd_title,
        source_jd,                               # optional JD filename (design D3)
        record.get("name"),
        record.get("title"),
        record.get("company"),
        record.get("location"),
        record.get("track"),
        record.get("profile_url"),
        ";".join(record.get("source_urls") or []),
        record.get("score"),
        record.get("justification"),
        "new",                                   # always initial status (spec)
        bool(record.get("dedup_status") == "possible-duplicate"),
        "",
    ]


# ------------------------------------------------------------------- reading

def load_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []
    ws = load_workbook(path, read_only=True).active
    values = ws.iter_rows(values_only=True)
    header = [str(c or "") for c in next(values)]
    rows = []
    for raw in values:
        if all(v is None for v in raw):
            continue
        rows.append(dict(zip(header, raw)))
    return rows


def excluded_urls(path: Path) -> set[str]:
    """Profile/source URLs of candidates already seen/rejected etc.

    Spec: only status 'new' is eligible for re-shortlisting consideration;
    terminal statuses (rejected/contacted/interviewing/hired) exclude forever.
    """
    out: set[str] = set()
    for row in load_rows(path):
        status = str(row.get("status") or "").strip().lower()
        if status in TERMINAL_STATUSES:
            for url in str(row.get("profile_url") or "").split(";"):
                if url:
                    out.add(url.strip())
            for url in str(row.get("source_urls") or "").split(";"):
                if url:
                    out.add(url.strip())
    return out


def prior_own_urls(path: Path, jd_key: str, jd_title: str | None = None,
                   terminal_statuses: set[str] | None = None) -> set[str]:
    """URLs of THIS JD's own prior non-terminal ledger rows (design D3).

    Used for re-added-JD reprocessing: these rows stop suppressing new
    selections. Matching is by `source_jd` (filename) when present; legacy
    rows without it fall back to normalized jd_title equality. Terminal
    reviewer statuses NEVER match - they always suppress unconditionally.
    """
    terminal = terminal_statuses or TERMINAL_STATUSES

    def _norm(value: str | None) -> str:
        return " ".join(str(value or "").lower().split())

    want_title = _norm(jd_title)
    out: set[str] = set()
    for row in load_rows(path):
        status = str(row.get("status") or "").strip().lower()
        if status in terminal:
            continue
        row_jd = str(row.get("source_jd") or "").strip()
        if row_jd:
            if row_jd != jd_key:
                continue
        elif not want_title or _norm(row.get("jd_title")) != want_title:
            continue                            # legacy row: title must match
        for url in str(row.get("profile_url") or "").split(";"):
            if url.strip():
                out.add(url.strip())
        for url in str(row.get("source_urls") or "").split(";"):
            if url.strip():
                out.add(url.strip())
    return out


# ------------------------------------------------------------------- writing

def append_records(path: Path, records: list[dict], jd_title: str,
                   backup: bool = True,
                   source_jd: str | None = None) -> int:
    """Append new rows (status 'new'); never modifies existing rows.

    `source_jd` (optional JD filename) enables reprocessing-aware dedup for
    future runs (design D3). Raises LedgerLockedError - WITHOUT touching the
    file - when the workbook seems open in Excel; callers defer to next run.
    """
    if _lock_files(path):
        raise LedgerLockedError(f"workbook looks open in Excel: {path.name}")
    ensure_workbook(path)

    if backup:
        shutil.copy2(path, path.with_suffix(".backup.xlsx"))

    try:
        wb = load_workbook(path)
    except PermissionError as exc:
        raise LedgerLockedError(str(exc)) from exc

    ws = wb.active
    existing_before = ws.max_row
    for record in records:
        ws.append(_record_to_row(record, jd_title, source_jd))
    try:
        wb.save(path)
    except PermissionError as exc:
        raise LedgerLockedError(str(exc)) from exc
    return ws.max_row - existing_before


def restore_backup(path: Path) -> bool:
    backup = path.with_suffix(".backup.xlsx")
    if not backup.exists():
        return False
    shutil.copy2(backup, path)
    return True
