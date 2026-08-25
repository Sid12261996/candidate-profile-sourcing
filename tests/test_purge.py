"""Tests for purge.py - retention window, dry-run vs apply (task 10.3)."""
import importlib
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "skill" / "candidate-sourcing" / "scripts"
sys.path.insert(0, str(SCRIPTS))

purge = importlib.import_module("purge")

CFG = {"retention": {"purge_rejected_after_months": 6}}

COLUMNS = ["date_added", "jd_title", "candidate_name", "status"]
TODAY = date(2026, 8, 25)
OLD = TODAY - timedelta(days=200)
RECENT = TODAY - timedelta(days=30)


def seed(tmp_path):
    path = tmp_path / "candidates.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(COLUMNS)
    ws.append([OLD.isoformat(), "JD A", "Old Rejected", "rejected"])
    ws.append([RECENT.isoformat(), "JD A", "Recent Rejected", "rejected"])
    ws.append([OLD.isoformat(), "JD A", "Old Contacted", "contacted"])   # never purge
    ws.append([OLD.isoformat(), "JD B", "Old New", "new"])               # never purge
    wb.save(path)
    return path


def test_plan_targets_only_old_rejected(tmp_path):
    path = seed(tmp_path)
    doomed = purge.plan_purge(path, CFG, today=TODAY)
    assert doomed == [2]          # worksheet row of Old Rejected only


def test_dry_run_default_changes_nothing(tmp_path):
    path = seed(tmp_path)
    before = path.read_bytes()
    purge.main([str(path)])       # no --apply
    assert path.read_bytes() == before


def test_apply_removes_only_planned_rows(tmp_path):
    path = seed(tmp_path)
    purge.apply_purge(path, purge.plan_purge(path, CFG, today=TODAY))
    names = [r[2] for r in Workbook.__dict__ and __import__("openpyxl").load_workbook(path).active.iter_rows(min_row=2, values_only=True)]
    assert names == ["Recent Rejected", "Old Contacted", "Old New"]
