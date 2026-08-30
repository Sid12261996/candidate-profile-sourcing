"""Tests for ledger.py - schema, append-only, exclusions, guardrails (tasks 8.1-8.2)."""
import importlib
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "skill" / "candidate-sourcing" / "scripts"
sys.path.insert(0, str(SCRIPTS))

ledger = importlib.import_module("ledger")


@pytest.fixture
def wb_path(tmp_path):
    return tmp_path / "candidates.xlsx"


def rec(name="Priya", score=90, url="u://priya"):
    return {"name": name, "title": "Designer", "company": "StudioX",
            "location": "Mumbai", "track": "xray", "profile_url": url,
            "source_urls": [url], "score": score,
            "justification": "strong portfolio"}


# ------------------------------------------------------------------- 8.1

def test_fresh_workbook_created_with_schema(wb_path):
    assert not wb_path.exists()
    ledger.append_records(wb_path, [rec()], "Lead Designer")
    rows = ledger.load_rows(wb_path)
    assert list(rows[0].keys()) == ledger.LEDGER_COLUMNS


def test_append_only_new_rows_status_new_and_no_mutation(wb_path):
    ledger.append_records(wb_path, [rec(name="First")], "JD A")
    before = ledger.load_rows(wb_path)[0]

    ledger.append_records(wb_path, [rec(name="Second", url="u://second")], "JD B")
    rows = ledger.load_rows(wb_path)
    assert len(rows) == 2
    after = [r for r in rows if r["candidate_name"] == "First"][0]
    assert dict(before) == dict(after)          # existing row untouched
    new = [r for r in rows if r["candidate_name"] == "Second"][0]
    assert new["status"] == "new"
    assert new["date_added"]                    # dated today


def test_terminal_statuses_drive_exclusions(wb_path):
    ledger.append_records(wb_path, [
        rec(name="A", url="u://a"), rec(name="R", url="u://r"),
        rec(name="C", url="u://c"), rec(name="H", url="u://h"),
        rec(name="N", url="u://n"),
    ], "JD")
    ws = Workbook.__init__  # noqa - readability only; statuses set via direct edit below

    from openpyxl import load_workbook
    book = load_workbook(wb_path)
    sheet = book.active
    status_col = ledger.LEDGER_COLUMNS.index("status") + 1
    for row_idx, status in [(2, "rejected"), (3, "contacted"),
                            (4, "interviewing"), (5, "hired"), (6, "new")]:
        sheet.cell(row=row_idx, column=status_col, value=status)
    book.save(wb_path)

    excluded = ledger.excluded_urls(wb_path)
    assert excluded == {"u://a", "u://r", "u://c", "u://h"}   # 'new' NOT excluded


# ------------------------------------------------ prior_own_urls / source_jd (4.1)

def test_source_jd_column_written_when_provided(wb_path):
    ledger.append_records(wb_path, [rec()], "Workspace Designer",
                          source_jd="zyeta-role.md")
    row = ledger.load_rows(wb_path)[0]
    assert row["source_jd"] == "zyeta-role.md"


def test_prior_own_urls_returns_own_non_terminal_only(wb_path):
    ledger.append_records(wb_path, [
        rec(name="OwnNew", url="u://own-new"),
        rec(name="OwnRej", url="u://own-rej"),
        rec(name="Other", url="u://other"),
    ], "Designer", source_jd="zyeta-role.md")

    from openpyxl import load_workbook
    book = load_workbook(wb_path)
    sheet = book.active
    status_col = ledger.LEDGER_COLUMNS.index("status") + 1
    jd_col = ledger.LEDGER_COLUMNS.index("source_jd") + 1
    sheet.cell(row=3, column=status_col, value="rejected")   # own but terminal
    sheet.cell(row=4, column=jd_col, value="some-other-jd.md")
    book.save(wb_path)

    own = ledger.prior_own_urls(wb_path, "zyeta-role.md")
    assert own == {"u://own-new"}          # terminal own row + other JD excluded


def test_prior_own_urls_legacy_title_fallback(wb_path):
    # legacy rows: no source_jd column value -> normalized title match
    ledger.append_records(wb_path, [rec(name="Old", url="u://old")], "Lead Designer")
    own = ledger.prior_own_urls(wb_path, "lead-designer.md", jd_title="lead designer")
    assert own == {"u://old"}
    # no title supplied -> no legacy match (never over-match)
    assert ledger.prior_own_urls(wb_path, "lead-designer.md") == set()


def test_prior_own_urls_empty_for_unknown_jd(wb_path):
    ledger.append_records(wb_path, [rec()], "Designer", source_jd="a.md")
    assert ledger.prior_own_urls(wb_path, "b.md", jd_title="Designer") == set()


# ------------------------------------------------------------------- 8.2

def test_backup_created_before_write(wb_path):
    ledger.append_records(wb_path, [rec()], "JD")
    backup = wb_path.with_suffix(".backup.xlsx")
    assert backup.exists()
    ledger.append_records(wb_path, [rec(name="B", url="u://b")], "JD")
    rows = ledger.load_rows(backup)
    assert len(rows) == 1                        # snapshot from BEFORE second write


def test_locked_file_defers_without_touching(wb_path):
    ledger.ensure_workbook(wb_path)              # simulate Excel lock file
    lock = wb_path.parent / f"~${wb_path.name}"
    lock.write_bytes(b"excel lock")

    with pytest.raises(ledger.LedgerLockedError):
        ledger.append_records(wb_path, [rec()], "JD")

    assert ledger.load_rows(wb_path) == []       # nothing appended


def test_restore_backup_recovers(wb_path):
    ledger.append_records(wb_path, [rec(name="Only")], "JD")
    ledger.append_records(wb_path, [rec(name="Second", url="u://second")], "JD")
    # backup now holds the post-first-write state (1 row)
    wb_path.write_bytes(b"corrupt")              # simulated bad write
    assert ledger.restore_backup(wb_path)
    rows = ledger.load_rows(wb_path)
    assert len(rows) == 1                        # recovered to pre-second-write state
    assert rows[0]["candidate_name"] == "Only"
